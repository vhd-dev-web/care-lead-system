from __future__ import annotations

import csv
from pathlib import Path


def read_table(path: str | Path) -> list[dict[str, str]]:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(file_path)
    return read_csv(file_path)


def read_csv(path: str | Path) -> list[dict[str, str]]:
    file_path = Path(path)
    text = file_path.read_text(encoding="utf-8-sig")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    with file_path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle, dialect=dialect)
        return [dict(row) for row in reader]


def read_xlsx(path: str | Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = ["" if value is None else str(value).strip() for value in rows[0]]
    result: list[dict[str, str]] = []
    for values in rows[1:]:
        row = {
            headers[index]: "" if value is None else str(value).strip()
            for index, value in enumerate(values)
            if index < len(headers) and headers[index]
        }
        if any(row.values()):
            result.append(row)
    return result
