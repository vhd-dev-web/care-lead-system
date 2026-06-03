from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .master_schema import MASTER_COLUMNS, ordered_fieldnames
from .master_store import MasterStore
from .run_registry import RunRegistry


def export_master_xlsx(root_dir: str | Path = ".", parent_run_id: str = "") -> Path:
    store = MasterStore(root_dir)
    rows = store.load_rows()
    fieldnames = ordered_fieldnames(rows) if rows else MASTER_COLUMNS
    store.ensure_dirs()

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Leads Master"
    worksheet.append(fieldnames)
    for row in rows:
        worksheet.append([row.get(field, "") for field in fieldnames])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    worksheet.freeze_panes = "A2"
    if fieldnames:
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(fieldnames))}{max(1, len(rows) + 1)}"
    for index, field in enumerate(fieldnames, start=1):
        max_len = max([len(field)] + [len(str(row.get(field, ""))) for row in rows[:200]])
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(max_len + 2, 12), 42)

    try:
        workbook.save(store.master_xlsx_path)
        save_status = "done"
        save_error = ""
    except PermissionError as exc:
        # The .xlsx is a human-readable mirror of the CSV. If it is open
        # in Excel or syncing in Drive, the file lock prevents us from
        # overwriting it. Do not let that abort the rest of the pipeline
        # (clay queue export and google sheet sync still need to run).
        save_status = "skipped"
        save_error = f"PermissionError: {exc}. XLSX file is locked (open in Excel or Drive sync?). CSV master is unaffected."
        print(f"[export_master_xlsx] WARNING: {save_error}")

    if not parent_run_id:
        registry = RunRegistry(root_dir)
        run = registry.start_run("export_master_xlsx")
        registry.finish_run(
            run.run_id,
            save_status,
            output_paths=[str(store.master_xlsx_path)] if save_status == "done" else [],
            records_read=len(rows),
            records_written=len(rows) if save_status == "done" else 0,
            error=save_error,
        )
    return store.master_xlsx_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Generate data/master/Leads_Master.xlsx from the master CSV.")
    parser.add_argument("--root", default=".", help="Repository/workspace root.")
    args = parser.parse_args(argv)
    path = export_master_xlsx(root_dir=args.root)
    print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
